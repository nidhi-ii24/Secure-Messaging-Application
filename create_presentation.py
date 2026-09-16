"""
Generates the PowerPoint presentation for CNS Secure Messaging System.
Focuses strictly on:
1. Problem Identification
2. System & Cryptographic Design
3. Implementation Planning
(Excludes any mention of completion).
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_filename="CNS_Secure_Messaging_Proposal.pptx"):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank layout

    # Palette
    COLOR_BG = RGBColor(15, 23, 42)        # #0F172A (Deep Slate)
    COLOR_CARD = RGBColor(30, 41, 59)      # #1E293B (Card Slate)
    COLOR_CARD_BORDER = RGBColor(51, 65, 85) # #334155
    COLOR_ACCENT = RGBColor(56, 189, 248)  # #38BDF8 (Cyan Accent)
    COLOR_ACCENT_ALT = RGBColor(2, 132, 199) # #0284C7
    COLOR_TEXT_MAIN = RGBColor(248, 250, 252) # #F8FAFC
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8
    COLOR_TAG_BG = RGBColor(12, 74, 110)
    COLOR_RED = RGBColor(239, 68, 68)
    COLOR_GREEN = RGBColor(34, 197, 94)

    def apply_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background() # no line

    def add_header(slide, title_text, category_text="CNS PROJECT PROPOSAL"):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.35))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category_text.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_ACCENT

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_MAIN

    def add_card(slide, left, top, width, height, title=None, subtitle=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1)

        offset_top = top + Inches(0.2)
        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.25), offset_top, width - Inches(0.5), Inches(0.4))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(15)
            p.font.bold = True
            p.font.color.rgb = COLOR_ACCENT
            offset_top += Inches(0.35)

        if subtitle:
            tb_sub = slide.shapes.add_textbox(left + Inches(0.25), offset_top, width - Inches(0.5), Inches(0.3))
            tf_sub = tb_sub.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle
            p_sub.font.size = Pt(11)
            p_sub.font.color.rgb = COLOR_TEXT_MUTED
            offset_top += Inches(0.3)

        return offset_top

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_background(s1)

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.15), Inches(3.6))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_ACCENT
    bar.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.7), Inches(11.3), Inches(3.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "CRYPTOGRAPHY & NETWORK SECURITY  |  PROJECT PROPOSAL"
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_ACCENT

    p_main = tf1.add_paragraph()
    p_main.text = "Secure End-to-End Messaging System\nUsing Hybrid Cryptography"
    p_main.font.size = Pt(36)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_TEXT_MAIN
    p_main.space_before = Pt(14)
    p_main.space_after = Pt(14)

    p_sub = tf1.add_paragraph()
    p_sub.text = "Implementation Design of Diffie–Hellman Key Exchange, AES Encryption, HMAC-Based Integrity, and TLS-Secured Communication"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED

    p_scope = tf1.add_paragraph()
    p_scope.text = "Focus Areas: Problem Identification  •  System & Cryptographic Design  •  Implementation Planning"
    p_scope.font.size = Pt(13)
    p_scope.font.bold = True
    p_scope.font.color.rgb = COLOR_ACCENT
    p_scope.space_before = Pt(20)

    # ==========================================
    # SLIDE 2: PRESENTATION OUTLINE
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_background(s2)
    add_header(s2, "Presentation Structure", "Agenda Overview")

    agenda_items = [
        ("01", "Problem Identification", "Threat analysis of conventional message relay systems, risks of untrusted networks, and formal problem definition."),
        ("02", "System & Security Design", "End-to-end architecture, hybrid cryptographic pipeline, key derivation, and defense mechanisms."),
        ("03", "Implementation Planning", "Modular work breakdown structure (WBS), execution roadmap, and security verification strategy.")
    ]

    for i, (num, heading, desc) in enumerate(agenda_items):
        top_pos = Inches(1.8 + i * 1.6)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.7), Inches(1.3))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CARD_BORDER

        # Number Badge
        nb = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), top_pos + Inches(0.25), Inches(0.8), Inches(0.8))
        nb.fill.solid()
        nb.fill.fore_color.rgb = COLOR_TAG_BG
        nb.line.color.rgb = COLOR_ACCENT
        tf_nb = nb.text_frame
        p_nb = tf_nb.paragraphs[0]
        p_nb.text = num
        p_nb.font.size = Pt(20)
        p_nb.font.bold = True
        p_nb.font.color.rgb = COLOR_ACCENT
        p_nb.alignment = PP_ALIGN.CENTER

        # Text
        tb = s2.shapes.add_textbox(Inches(2.2), top_pos + Inches(0.15), Inches(10.0), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = heading
        p_h.font.size = Pt(18)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_TEXT_MAIN

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(4)

    # ==========================================
    # SLIDE 3: PROBLEM IDENTIFICATION - CORE CHALLENGES
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_background(s3)
    add_header(s3, "Problem Identification: Insecure Message Transmission", "Problem Identification")

    # Main Problem Statement Box
    prob_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(1.6))
    prob_box.fill.solid()
    prob_box.fill.fore_color.rgb = COLOR_CARD
    prob_box.line.color.rgb = COLOR_RED
    prob_box.line.width = Pt(1.5)

    tf_pb = prob_box.text_frame
    tf_pb.word_wrap = True
    p_pbt = tf_pb.paragraphs[0]
    p_pbt.text = "PROBLEM STATEMENT"
    p_pbt.font.size = Pt(12)
    p_pbt.font.bold = True
    p_pbt.font.color.rgb = COLOR_RED

    p_pbc = tf_pb.add_paragraph()
    p_pbc.text = (
        "Conventional messaging applications operating across untrusted networks expose communication to eavesdropping, "
        "message tampering, unauthorized server access, and replay attacks. When messages are relayed or stored in plaintext "
        "at intermediary servers, an attacker or compromised host can access sensitive records or silently alter data without detection."
    )
    p_pbc.font.size = Pt(13.5)
    p_pbc.font.color.rgb = COLOR_TEXT_MAIN
    p_pbc.space_before = Pt(4)

    # 3 Specific Core Weaknesses
    weaknesses = [
        ("Intermediary Server Exposure", "Centralized relays traditionally terminate transport connections, exposing plaintext messages and credentials in server databases and memory buffers."),
        ("Key Exchange Vulnerability", "Transmitting symmetric keys directly over networks allows interceptors to obtain keys. Basic unauthenticated DH remains vulnerable to active MITM substitution."),
        ("Lack of Independent Integrity", "Without cryptographically binding IV, ciphertext, and sequence nonces via MAC, adversaries can alter bits or replay legitimate intercepted commands.")
    ]

    for i, (title, text) in enumerate(weaknesses):
        left_pos = Inches(0.8 + i * 4.0)
        c_top = add_card(s3, left_pos, Inches(3.6), Inches(3.7), Inches(3.3), title=title)
        tb = s3.shapes.add_textbox(left_pos + Inches(0.2), c_top, Inches(3.3), Inches(2.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 4: PROBLEM IDENTIFICATION - THREAT MODEL & ATTACK VECTORS
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_background(s4)
    add_header(s4, "Threat Model & Specific Attack Vectors", "Problem Identification")

    attacks = [
        ("Attack 1: Eavesdropping / Sniffing", "Passive Attack", "Adversaries intercept unencrypted network packets (e.g. using packet analyzers like Wireshark) on untrusted LANs/Wi-Fi, reading plaintext messages and sensitive user data."),
        ("Attack 2: Message Tampering", "Active Attack", "An adversary in transit modifies bits in the ciphertext or parameters, attempting to alter transaction amounts, names, or instructions without recipient detection."),
        ("Attack 3: Replay Attack", "Active Attack", "An attacker captures a valid encrypted transaction packet and re-transmits it later, causing unintended duplicate actions (e.g. repeated funds transfer)."),
        ("Attack 4: Man-in-the-Middle (MITM)", "Asymmetric Attack", "In unauthenticated Diffie-Hellman, Eve intercepts public value A and replaces it with A', establishing separate shared secrets with both Alice and Bob while impersonating them.")
    ]

    for i, (title, tag, desc) in enumerate(attacks):
        row = i // 2
        col = i % 2
        left_pos = Inches(0.8 + col * 5.95)
        top_pos = Inches(1.8 + row * 2.6)

        c_top = add_card(s4, left_pos, top_pos, Inches(5.75), Inches(2.35), title=title, subtitle=tag)
        tb = s4.shapes.add_textbox(left_pos + Inches(0.2), c_top, Inches(5.35), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 5: SYSTEM DESIGN - ARCHITECTURAL OVERVIEW
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_background(s5)
    add_header(s5, "Proposed System Architecture & Zero-Knowledge Relay", "System Design")

    # 3 Layers
    layers = [
        ("1. Client Layer (Alice & Bob)", "End-to-End Cryptographic Boundary", [
            "Generates ephemeral Diffie-Hellman keypairs locally in client memory.",
            "Computes shared secret and derives symmetric keys via HKDF-SHA256.",
            "Performs client-side AES-256-CBC encryption and HMAC-SHA256 generation.",
            "Validates HMAC and decrypts messages strictly on client device."
        ]),
        ("2. Transport Layer (TLS / HTTPS)", "Communication Channel Security", [
            "Encapsulates client-server communication inside TLS 1.3 encrypted tunnels.",
            "Utilizes X.509 certificates to protect transport links against eavesdropping.",
            "Provides hop-by-hop transport confidentiality and channel integrity."
        ]),
        ("3. Server Relay Layer (Flask)", "Untrusted Relay & Message Broker", [
            "Coordinates public DH parameters (p, g, A, B) without learning private keys.",
            "Enforces anti-replay verification (message ID & sequence check).",
            "Stores & queues ONLY ciphertext, IV, HMAC tag, and timestamp.",
            "Core Security Invariant: Zero plaintext messages or session keys on server."
        ])
    ]

    for i, (title, sub, bullets) in enumerate(layers):
        left_pos = Inches(0.8 + i * 4.0)
        c_top = add_card(s5, left_pos, Inches(1.8), Inches(3.7), Inches(5.1), title=title, subtitle=sub)
        tb = s5.shapes.add_textbox(left_pos + Inches(0.2), c_top, Inches(3.3), Inches(3.5))
        tf = tb.text_frame
        tf.word_wrap = True

        for j, b in enumerate(bullets):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = f"• {b}"
            p.font.size = Pt(12)
            p.font.color.rgb = COLOR_TEXT_MUTED
            p.space_after = Pt(8)

    # ==========================================
    # SLIDE 6: SYSTEM DESIGN - HYBRID CRYPTOGRAPHIC PIPELINE
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_background(s6)
    add_header(s6, "Hybrid Cryptographic Pipeline Specification", "System Design")

    stages = [
        ("Key Exchange (DH)", "RFC 3526 MODP-2048", "Alice: a, A = g^a mod p\nBob: b, B = g^b mod p\nShared Secret S = g^(ab) mod p\nSecret never sent over network."),
        ("Key Derivation (HKDF)", "HKDF-SHA256 (RFC 5869)", "Derives independent keys:\n• K_enc = HKDF(S, 'aes-key')\n• K_mac = HKDF(S, 'hmac-key')\nSeparates confidentiality & integrity."),
        ("Confidentiality (AES)", "AES-256-CBC + PKCS#7", "Plaintext encrypted under K_enc.\nFresh 16-byte random IV per message.\nGuarantees message confidentiality against intermediaries."),
        ("Integrity & Auth (HMAC)", "HMAC-SHA256 (Encrypt-then-MAC)", "Tag = HMAC(IV || Ciphertext || TS || Seq, K_mac).\nRecipient verifies tag before decryption.\nRejects any altered bits.")
    ]

    for i, (title, sub, desc) in enumerate(stages):
        left_pos = Inches(0.8 + i * 3.0)
        c_top = add_card(s6, left_pos, Inches(1.8), Inches(2.75), Inches(5.1), title=title, subtitle=sub)
        tb = s6.shapes.add_textbox(left_pos + Inches(0.15), c_top, Inches(2.45), Inches(3.6))
        tf = tb.text_frame
        tf.word_wrap = True
        for line_idx, line in enumerate(desc.split("\n")):
            p = tf.paragraphs[0] if line_idx == 0 else tf.add_paragraph()
            p.text = line
            p.font.size = Pt(12)
            p.font.color.rgb = COLOR_TEXT_MAIN if line.startswith("•") else COLOR_TEXT_MUTED
            p.space_after = Pt(4)

    # ==========================================
    # SLIDE 7: SYSTEM DESIGN - PROTOCOL MESSAGE FLOW
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_background(s7)
    add_header(s7, "Step-by-Step Message Transmission Protocol", "System Design")

    steps = [
        ("Step 1: Ephemeral Handshake", "Alice & Bob exchange public values A and B via relay; verify RSA signatures to prevent MITM."),
        ("Step 2: Key Derivation", "Both compute identical shared secret S = g^(ab) mod p, then derive AES Key (K_enc) and HMAC Key (K_mac)."),
        ("Step 3: Encrypt-then-MAC", "Alice encrypts plaintext with AES-256-CBC (fresh IV), generates timestamp & seq#, computes HMAC-SHA256 tag."),
        ("Step 4: Untrusted Relay", "Server receives {sender, receiver, IV, Ciphertext, Tag, TS, Seq#}. Server checks replay filter and queues payload."),
        ("Step 5: Integrity Verification", "Bob recomputes HMAC over received ciphertext & metadata. If tag does NOT match, message is immediately rejected."),
        ("Step 6: Plaintext Recovery", "Once integrity is confirmed, Bob checks sequence freshness and decrypts ciphertext using AES-256-CBC.")
    ]

    for i, (title, desc) in enumerate(steps):
        row = i // 3
        col = i % 3
        left_pos = Inches(0.8 + col * 4.0)
        top_pos = Inches(1.8 + row * 2.6)

        c_top = add_card(s7, left_pos, top_pos, Inches(3.75), Inches(2.35), title=title)
        tb = s7.shapes.add_textbox(left_pos + Inches(0.2), c_top, Inches(3.35), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(12.5)
        p.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 8: SYSTEM DESIGN - PLANNED DEFENSE STRATEGY
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_background(s8)
    add_header(s8, "Design Mapping: Attacks vs Proposed Defenses", "System Design")

    matrix = [
        ("Attack Vector", "Vulnerability Cause", "Proposed Cryptographic Defense", "Expected Security Guarantee"),
        ("Eavesdropping", "Plaintext transmission over network", "Dual-Layer: TLS 1.3 (Transport) + AES-256-CBC (E2EE)", "Zero plaintext leakage to network sniffers or relay server"),
        ("Message Tampering", "Adversary alters bits in transit", "HMAC-SHA256 with Encrypt-then-MAC & constant-time check", "Altered ciphertext causes tag mismatch; discarded before decrypt"),
        ("Replay Attack", "Retransmission of valid intercepted packet", "Unique UUID message_id, monotonic seq#, and ISO timestamp", "Duplicate packet IDs rejected by server & client filters (HTTP 409)"),
        ("DH Man-in-the-Middle", "Unauthenticated public parameter exchange", "RSA-2048 PSS Digital Signatures on public DH parameters", "Forged public keys fail signature check; connection aborted"),
        ("Credential Leaks", "Plaintext or reversible password storage", "PBKDF2-HMAC-SHA256 (100k rounds) with random 16B salt", "Resistant to rainbow table and GPU dictionary attacks")
    ]

    table_shape = s8.shapes.add_table(6, 4, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(2.8)
    table.columns[2].width = Inches(3.6)
    table.columns[3].width = Inches(3.1)

    for r_idx, row in enumerate(matrix):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            cell.fill.solid()
            if r_idx == 0:
                cell.fill.fore_color.rgb = COLOR_TAG_BG
            else:
                cell.fill.fore_color.rgb = COLOR_CARD if r_idx % 2 == 0 else RGBColor(24, 34, 52)
            
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11.5 if r_idx > 0 else 12)
            p.font.bold = (r_idx == 0 or c_idx == 0)
            p.font.color.rgb = COLOR_ACCENT if r_idx == 0 else (COLOR_TEXT_MAIN if c_idx == 2 else COLOR_TEXT_MUTED)

    # ==========================================
    # SLIDE 9: PLANNING - WORK BREAKDOWN STRUCTURE (WBS)
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_background(s9)
    add_header(s9, "Implementation Planning: Work Breakdown Structure", "Implementation Planning")

    phases = [
        ("Phase 1: Crypto Engine Design", "Core Primitives", [
            "Specification of RFC 3526 MODP-2048 Diffie-Hellman algorithm.",
            "HKDF-SHA256 key separation logic (AES and HMAC keys).",
            "AES-256-CBC encryption/decryption module with PKCS#7 padding.",
            "HMAC-SHA256 generation and constant-time comparison logic.",
            "RSA-2048 PSS digital signature module for MITM prevention."
        ]),
        ("Phase 2: Server & Database Design", "Backend Architecture", [
            "Relational SQLite schema design (users, dh_sessions, messages).",
            "Authentication endpoints with PBKDF2-HMAC-SHA256 password hashing.",
            "Public DH parameter exchange coordination routes.",
            "Encrypted message queuing and anti-replay filter implementation.",
            "Zero-plaintext database design verification."
        ]),
        ("Phase 3: Client & UI Architecture", "Client Application", [
            "Client-side cryptographic manager for local key isolation.",
            "Dual-user communication interface (Alice and Bob simulation).",
            "Real-time Cryptographic Pipeline Inspector for viva/demo visibility.",
            "Simulated adversary controls (in-flight tampering injection)."
        ])
    ]

    for i, (title, sub, bullets) in enumerate(phases):
        left_pos = Inches(0.8 + i * 4.0)
        c_top = add_card(s9, left_pos, Inches(1.8), Inches(3.7), Inches(5.1), title=title, subtitle=sub)
        tb = s9.shapes.add_textbox(left_pos + Inches(0.2), c_top, Inches(3.3), Inches(3.6))
        tf = tb.text_frame
        tf.word_wrap = True

        for j, b in enumerate(bullets):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = f"• {b}"
            p.font.size = Pt(12)
            p.font.color.rgb = COLOR_TEXT_MUTED
            p.space_after = Pt(7)

    # ==========================================
    # SLIDE 10: PLANNING - VERIFICATION & ATTACK TESTING PLAN
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    apply_background(s10)
    add_header(s10, "Planning: Security Verification & Attack Evaluation Plan", "Implementation Planning")

    test_plans = [
        ("Cryptographic Correctness Plan", "Unit Testing Suite", [
            "Verify mathematical agreement of Diffie-Hellman: S_A == S_B.",
            "Verify HKDF produces two strictly distinct 256-bit symmetric keys.",
            "Verify AES-256-CBC encryption/decryption round-trip with arbitrary text.",
            "Verify constant-time HMAC tag generation and verification.",
            "Verify RSA-2048 signature generation and rejection of forged signatures."
        ]),
        ("Attack Simulation Evaluation Plan", "Vulnerability Testing", [
            "Eavesdropping Test: Compare plaintext HTTP vs TLS+AES payload in network captures.",
            "Tampering Test: Programmatically alter 1 byte in ciphertext; verify HMAC triggers alert.",
            "Replay Test: Re-send legitimate packet ID; assert rejection with HTTP 409 Conflict.",
            "MITM Test: Simulate adversary key substitution; verify RSA signature catches mismatch."
        ]),
        ("Database Security Audit Plan", "Zero-Plaintext Verification", [
            "Query raw SQLite tables during active messaging.",
            "Assert zero plaintext message characters exist in database columns.",
            "Confirm server never accesses or holds symmetric encryption keys."
        ])
    ]

    for i, (title, sub, bullets) in enumerate(test_plans):
        left_pos = Inches(0.8 + i * 4.0)
        c_top = add_card(s10, left_pos, Inches(1.8), Inches(3.7), Inches(5.1), title=title, subtitle=sub)
        tb = s10.shapes.add_textbox(left_pos + Inches(0.2), c_top, Inches(3.3), Inches(3.6))
        tf = tb.text_frame
        tf.word_wrap = True

        for j, b in enumerate(bullets):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = f"• {b}"
            p.font.size = Pt(12)
            p.font.color.rgb = COLOR_TEXT_MUTED
            p.space_after = Pt(7)

    # ==========================================
    # SLIDE 11: PLANNING - TECHNOLOGY STACK & SYLLABUS MAPPING
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    apply_background(s11)
    add_header(s11, "Technology Selection & Course Outcome (CO) Alignment", "Implementation Planning")

    # Left Card: Tech Stack
    c_top1 = add_card(s11, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), title="Proposed Technology Stack", subtitle="Standard, Educational, & Lightweight")
    tb1 = s11.shapes.add_textbox(Inches(1.0), c_top1, Inches(5.3), Inches(3.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    tech_items = [
        ("Language & Backend", "Python 3.14 + Flask REST API"),
        ("Cryptographic Library", "Python cryptography (hazmat) + hashlib & hmac"),
        ("Database Layer", "SQLite3 (Lightweight, zero-config relational store)"),
        ("Transport Security", "TLS 1.3 / HTTPS with self-signed X.509 certificates"),
        ("User Interface", "HTML5, CSS3, Modern JavaScript (Dual-panel simulator)"),
        ("Testing Framework", "Python unittest (Automated security & crypto suites)")
    ]

    for idx, (k, v) in enumerate(tech_items):
        p = tf1.paragraphs[0] if idx == 0 else tf1.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(10)

    # Right Card: CO Mapping
    c_top2 = add_card(s11, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), title="Alignment with CNS Course Outcomes", subtitle="Curriculum Objectives & Competencies")
    tb2 = s11.shapes.add_textbox(Inches(7.0), c_top2, Inches(5.3), Inches(3.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    co_items = [
        ("CO1: Algorithms & Key Management", "Covered by AES-256, Diffie-Hellman MODP-2048, and HKDF-SHA256 key derivation."),
        ("CO2: Access & Security Mechanisms", "Covered by PBKDF2 password authentication, session tokens, and anti-replay nonces."),
        ("CO3: Secure Communication Protocols", "Covered by native TLS 1.3 / HTTPS transport and application-layer Encrypt-then-MAC."),
        ("CO4: Implementation & Evaluation", "Covered by complete prototype architecture, verification plan, and attack test suite.")
    ]

    for idx, (co, detail) in enumerate(co_items):
        p = tf2.paragraphs[0] if idx == 0 else tf2.add_paragraph()
        p.text = f"{co}\n  ➔ {detail}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(12)

    # ==========================================
    # SLIDE 12: CONCLUSION / Q&A
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    apply_background(s12)

    card_fin = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0), Inches(1.6), Inches(9.333), Inches(4.3))
    card_fin.fill.solid()
    card_fin.fill.fore_color.rgb = COLOR_CARD
    card_fin.line.color.rgb = COLOR_ACCENT
    card_fin.line.width = Pt(1.5)

    tb_fin = s12.shapes.add_textbox(Inches(2.3), Inches(1.9), Inches(8.733), Inches(3.7))
    tf_fin = tb_fin.text_frame
    tf_fin.word_wrap = True

    p_t = tf_fin.paragraphs[0]
    p_t.text = "Summary of Proposal"
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_ACCENT
    p_t.alignment = PP_ALIGN.CENTER

    summary_bullets = [
        "Problem Identification: Detailed vulnerabilities of plain/unauthenticated message relays across untrusted networks.",
        "System & Cryptographic Design: Robust hybrid pipeline combining Diffie-Hellman, HKDF, AES-256, HMAC-SHA256, RSA signatures, and TLS.",
        "Implementation Planning: Structured work breakdown structure, security verification matrix, and attack testing strategy.",
        "Zero-Plaintext Guarantee: Intermediaries strictly handle ciphertext and cryptographic tags."
    ]

    for b in summary_bullets:
        p = tf_fin.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_before = Pt(8)

    p_qa = tf_fin.add_paragraph()
    p_qa.text = "Thank You! Questions & Discussion"
    p_qa.font.size = Pt(20)
    p_qa.font.bold = True
    p_qa.font.color.rgb = COLOR_TEXT_MAIN
    p_qa.alignment = PP_ALIGN.CENTER
    p_qa.space_before = Pt(18)

    # Save presentation
    prs.save(output_filename)
    print(f"[+] Presentation saved successfully as: {output_filename}")

if __name__ == "__main__":
    create_deck()
